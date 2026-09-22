#!/usr/bin/env python
"""Build the team worklog from repository evidence, not from memory.

    python scripts/build_team_worklog.py

Writes ``reports/team_worklog.xlsx`` and the evidence tables inside
``docs/team/WORKLOG.md``.

What this can and cannot derive
-------------------------------
It can derive, mechanically:

* who authored each commit, and when, to the second;
* which components each commit touched, from the paths it changed;
* the **observed commit span** of a working session — the time between the first
  and last commit in a cluster;
* when each machine-generated evidence artifact was produced.

It cannot derive *effort*. Work happens before the first commit of a session and
between commits, and none of it is timestamped. A session with one commit has no
span at all. So the hours column is filled only where a span is observable, it is
labelled a **lower bound**, and the total is never presented as the engagement
duration.

That is the honest limit of Git as a timesheet, and the reason the 20-hour
requirement is closed by a team attestation with this as supporting evidence
rather than by this file alone.

Sessions are clusters of commits by the same author separated by less than
``SESSION_GAP_HOURS``. The gap is a convention, not a measurement, and is printed
in the output so a reader can judge it.
"""

from __future__ import annotations

import json
import subprocess
import sys
from collections import Counter
from datetime import datetime, timedelta, timezone
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[1]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

XLSX_OUT = REPO_ROOT / "reports" / "team_worklog.xlsx"
MD_OUT = REPO_ROOT / "docs" / "team" / "WORKLOG.md"

#: Commits by one author closer together than this are treated as one session.
SESSION_GAP_HOURS = 3.0

#: Path prefix -> the component a reader would recognise. First match wins, so
#: the more specific prefixes come first.
COMPONENTS: tuple[tuple[str, str], ...] = (
    ("synthetic_data/education", "Education synthetic data"),
    ("synthetic_data_education_loans", "Education synthetic data"),
    ("synthetic_data/mortgage", "Mortgage synthetic data"),
    ("synthetic_data", "Synthetic data"),
    ("src/rag", "Mortgage/Education policy RAG"),
    ("src/memory", "Memory"),
    ("src/guardrails", "Guardrails / security"),
    ("src/observability", "Observability / Phoenix"),
    ("src/mcp_host", "MCP"),
    ("mcp_server", "MCP"),
    ("src/rule_families", "Rule engine"),
    ("src/rules.py", "Rule engine"),
    ("src/calculations.py", "Rule engine"),
    ("src/review_triggers.py", "Rule engine"),
    ("src/graph.py", "LangGraph"),
    ("src/supervisor.py", "LangGraph"),
    ("src/context", "Context engineering"),
    ("src/api", "API / UI"),
    ("src/web", "API / UI"),
    ("src/cli.py", "API / UI"),
    ("src/tools", "Agent tools"),
    ("src", "Core runtime"),
    ("eval/agent", "DeepEval / agent evaluation"),
    ("eval/retrieval", "Retrieval evaluation"),
    ("eval", "Evaluation"),
    ("tests", "Tests"),
    ("requirements_validation_tests", "Requirement validation"),
    ("scripts", "Evidence regeneration"),
    ("docs/team", "Team evidence"),
    ("docs/rag", "Documentation"),
    ("docs", "Governance / documentation"),
    ("reports", "Reports / evidence"),
    ("traces", "Observability / Phoenix"),
    ("logs", "Reports / evidence"),
    ("data", "Indexes / data"),
    ("config", "Configuration"),
)


def _git(*args: str) -> str:
    out = subprocess.run(["git", *args], cwd=str(REPO_ROOT), capture_output=True,
                         text=True, encoding="utf-8", errors="replace", check=False)
    if out.returncode != 0:
        raise SystemExit(f"git {' '.join(args)} failed: {out.stderr[:300]}")
    return out.stdout


def component_of(path: str) -> str:
    p = path.replace("\\", "/")
    for prefix, name in COMPONENTS:
        if p.startswith(prefix):
            return name
    return "Other"


def load_commits() -> list[dict]:
    """Every commit with author, UTC timestamp, subject and touched components."""
    raw = _git("log", "--all", "--no-merges", "--date=iso-strict",
               "--pretty=format:%x01%H%x02%an%x02%ae%x02%aI%x02%s")
    commits: list[dict] = []
    for block in raw.split("\x01"):
        if not block.strip():
            continue
        sha, name, email, iso, subject = block.strip().split("\x02", 4)
        when = datetime.fromisoformat(iso).astimezone(timezone.utc)
        files = [f for f in _git("show", "--name-only", "--pretty=format:", sha)
                 .splitlines() if f.strip()]
        counts = Counter(component_of(f) for f in files)
        commits.append({
            "sha": sha[:7],
            "author": name,
            "email": email,
            "when": when,
            "subject": subject,
            "files": len(files),
            "components": [c for c, _ in counts.most_common(4)],
        })
    commits.sort(key=lambda c: c["when"])
    return commits


def build_sessions(commits: list[dict]) -> list[dict]:
    """Cluster each author's commits into working sessions."""
    by_author: dict[str, list[dict]] = {}
    for c in commits:
        by_author.setdefault(c["author"], []).append(c)

    sessions: list[dict] = []
    for author, rows in by_author.items():
        rows.sort(key=lambda c: c["when"])
        current: list[dict] = []
        for c in rows:
            if current and (c["when"] - current[-1]["when"]) > timedelta(hours=SESSION_GAP_HOURS):
                sessions.append(_session(author, current))
                current = []
            current.append(c)
        if current:
            sessions.append(_session(author, current))
    sessions.sort(key=lambda s: s["start"])
    return sessions


def _session(author: str, rows: list[dict]) -> dict:
    start, end = rows[0]["when"], rows[-1]["when"]
    span_h = round((end - start).total_seconds() / 3600.0, 2)
    comps: list[str] = []
    for c in rows:
        for comp in c["components"]:
            if comp not in comps:
                comps.append(comp)
    single = len(rows) == 1
    return {
        "author": author,
        "start": start,
        "end": end,
        "commits": [c["sha"] for c in rows],
        "subjects": [c["subject"] for c in rows],
        "files": sum(c["files"] for c in rows),
        "components": comps,
        "span_hours": None if single else span_h,
        "basis": (
            "Single commit - no span is observable between commits, so no lower "
            "bound can be derived. Effort is team-attested."
            if single else
            f"Observed span between first and last commit in the session "
            f"({rows[0]['sha']} to {rows[-1]['sha']}). Lower bound: excludes work "
            f"before the first commit."
        ),
    }


def evidence_artifacts() -> list[dict]:
    """Machine-generated artifacts and the timestamp each records for itself."""
    targets = [
        ("reports/eval_report.json", "DeepEval / agent evaluation", "generated_at_utc"),
        ("reports/golden_signals.json", "Observability / Phoenix", "generated_at_utc"),
        ("reports/requirements_judge.json", "Requirement validation", "generated_at"),
        ("traces/trace_summary.json", "Observability / Phoenix", "generated_at_utc"),
        ("eval/results/retrieval_eval.json", "Retrieval evaluation", "generated_at_utc"),
        ("eval/results/retrieval_ablation.json", "Retrieval evaluation", "generated_at_utc"),
        ("eval/results/pipeline_sweep.json", "Retrieval evaluation", "generated_at_utc"),
        ("requirements_validation_tests/reports/latest_test_report.json",
         "Requirement validation", "generated_at"),
    ]
    rows = []
    for rel, component, key in targets:
        p = REPO_ROOT / rel
        if not p.exists():
            continue
        try:
            data = json.loads(p.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            continue
        rows.append({
            "artifact": rel,
            "component": component,
            "generated_at": str(data.get(key) or "-"),
            "runtime_s": data.get("wall_clock_seconds") or data.get("duration_seconds") or "",
        })
    rows.sort(key=lambda r: r["generated_at"])
    return rows


ATTESTATION = (
    "Krishna Annavaram and Mahesh Rajendra attest that combined project "
    "development, testing, evaluation, review and documentation activity for "
    "CredPilot met or exceeded the stated 20-hour engagement duration."
)

LIMIT_NOTE = (
    "Exact task-level hours are team-attested rather than mechanically derivable "
    "from Git history."
)


# ------------------------------------------------------------------ spreadsheet


def write_xlsx(commits, sessions, artifacts, path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1F4E79")
    wrap = Alignment(vertical="top", wrap_text=True)
    top = Alignment(vertical="top")

    wb = Workbook()

    def sheet(title, headers, rows, widths):
        ws = wb.create_sheet(title) if wb.sheetnames != ["Sheet"] else wb.active
        ws.title = title
        ws.append(headers)
        for c in ws[1]:
            c.font, c.fill, c.alignment = head_font, head_fill, top
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = wrap
        ws.freeze_panes = "A2"
        return ws

    # -- Worklog ---------------------------------------------------------------
    rows = []
    for s in sessions:
        rows.append([
            s["start"].strftime("%Y-%m-%d"),
            s["author"],
            ", ".join(s["components"][:2]) or "Other",
            ", ".join(s["components"]) or "Other",
            "; ".join(s["subjects"])[:400],
            ", ".join(s["commits"]),
            f"{s['start'].strftime('%Y-%m-%d %H:%M')} - {s['end'].strftime('%H:%M')} UTC",
            s["span_hours"] if s["span_hours"] is not None else "not derivable",
            s["basis"],
            "",
        ])
    sheet("Worklog",
          ["Date", "Member", "Work area", "Component", "Activity",
           "Evidence reference", "Observed timestamp range",
           "Observed hours (lower bound)", "Estimate basis", "Reviewed/confirmed by"],
          rows, [12, 20, 26, 34, 52, 20, 30, 16, 46, 22])

    # -- Commits ---------------------------------------------------------------
    sheet("Commits",
          ["Commit", "Author", "Timestamp (UTC)", "Files changed", "Components", "Subject"],
          [[c["sha"], c["author"], c["when"].strftime("%Y-%m-%d %H:%M:%S"),
            c["files"], ", ".join(c["components"]), c["subject"]] for c in commits],
          [11, 20, 21, 14, 40, 60])

    # -- Machine evidence ------------------------------------------------------
    sheet("Machine evidence",
          ["Artifact", "Component", "Self-recorded generation time", "Runtime (s)"],
          [[a["artifact"], a["component"], a["generated_at"], a["runtime_s"]]
           for a in artifacts],
          [62, 32, 30, 14])

    # -- Attestation -----------------------------------------------------------
    ws = sheet("Attestation", ["Field", "Value"], [], [30, 96])
    for field, value in [
        ("Requirement", "REQ-008 - Duration | 20 hours"),
        ("Statement", ATTESTATION),
        ("Derivation limit", LIMIT_NOTE),
        ("Commit window (UTC)",
         f"{commits[0]['when'].strftime('%Y-%m-%d %H:%M')} to "
         f"{commits[-1]['when'].strftime('%Y-%m-%d %H:%M')}"),
        ("Elapsed wall-clock (hours)",
         round((commits[-1]["when"] - commits[0]["when"]).total_seconds() / 3600.0, 1)),
        ("Note on wall-clock",
         "Elapsed wall-clock is the window work happened inside. It is not a "
         "measure of effort and is not offered as one."),
        ("Session gap convention (hours)", SESSION_GAP_HOURS),
        ("", ""),
        ("Krishna Annavaram - confirmation", ""),
        ("Krishna Annavaram - date", ""),
        ("Mahesh Rajendra - confirmation", ""),
        ("Mahesh Rajendra - date", ""),
        ("", ""),
        ("Status", "PENDING - unsigned until both members complete their own row"),
    ]:
        ws.append([field, value])
    for row in ws.iter_rows(min_row=2):
        row[0].font = Font(bold=True)
        for c in row:
            c.alignment = wrap

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


# ------------------------------------------------------------------ markdown


def write_markdown(commits, sessions, artifacts, path: Path) -> None:
    total_observed = sum(s["span_hours"] or 0.0 for s in sessions)
    derivable = [s for s in sessions if s["span_hours"] is not None]
    window_h = round((commits[-1]["when"] - commits[0]["when"]).total_seconds() / 3600.0, 1)

    out = [
        "# CredPilot work log\n",
        "Generated by [`scripts/build_team_worklog.py`](../../scripts/build_team_worklog.py) ",
        "from Git history and the self-recorded timestamps inside the committed evidence ",
        "artifacts. Regenerate it with:\n",
        "```bash\npython scripts/build_team_worklog.py\n```\n",
        "Spreadsheet form: [`reports/team_worklog.xlsx`](../../reports/team_worklog.xlsx)\n",
        "\n## What this log is, and what it is not\n",
        f"> **{LIMIT_NOTE}**\n",
        "Git is not a timesheet. It records the moment a change was committed, not the "
        "hours spent arriving at it. Work before the first commit of a session, thinking "
        "between commits, reading, design discussion and review leave no timestamp at "
        "all — and on this team the majority of the research was done by Mahesh "
        "Rajendra and is not represented in the commit history for exactly that reason.\n",
        "So this log reports two different things and never mixes them:\n",
        "| | Derivable from the repository? |",
        "|---|---|",
        "| Commit authorship and timestamps | **Yes**, to the second |",
        "| Components touched per session | **Yes**, from changed paths |",
        "| Observed span between commits in a session | **Yes** — a *lower bound* |",
        "| Hours actually worked | **No** — team-attested |",
        "",
        "\n## Observable window\n",
        "```",
        f"First commit        : {commits[0]['sha']}  {commits[0]['when'].strftime('%Y-%m-%d %H:%M:%S')} UTC",
        f"Last commit         : {commits[-1]['sha']}  {commits[-1]['when'].strftime('%Y-%m-%d %H:%M:%S')} UTC",
        f"Elapsed wall-clock  : {window_h} hours",
        f"Commits (no merges) : {len(commits)}",
        f"Working sessions    : {len(sessions)}  (gap convention: {SESSION_GAP_HOURS} h)",
        f"Observed commit span: {round(total_observed, 2)} hours across "
        f"{len(derivable)} of {len(sessions)} sessions",
        "```\n",
        "**Elapsed wall-clock is not effort.** It is the window the work happened inside. "
        "**Observed commit span is a lower bound** on active work: it starts at the first "
        "commit of a session, so everything done before that commit is invisible to it. "
        "Neither figure is offered as the engagement duration.\n",
        "\n## Sessions\n",
        "| Date | Member | Components | Evidence | Observed range (UTC) | Observed h (lower bound) |",
        "| --- | --- | --- | --- | --- | --- |",
    ]
    for s in sessions:
        span = s["span_hours"] if s["span_hours"] is not None else "not derivable"
        out.append(
            f"| {s['start'].strftime('%Y-%m-%d')} | {s['author']} | "
            f"{', '.join(s['components'][:3]) or 'Other'} | "
            f"{', '.join(f'`{c}`' for c in s['commits'])} | "
            f"{s['start'].strftime('%H:%M')}–{s['end'].strftime('%H:%M')} | {span} |"
        )

    out += [
        "\n## Machine-generated evidence, with the time each records for itself\n",
        "These are independent of Git: each artifact was written by committed code and "
        "carries its own generation timestamp.\n",
        "| Artifact | Component | Generated | Runtime (s) |",
        "| --- | --- | --- | --- |",
    ]
    for a in artifacts:
        out.append(f"| `{a['artifact']}` | {a['component']} | {a['generated_at']} | "
                   f"{a['runtime_s'] or '—'} |")

    out += [
        "\n## Work not represented above\n",
        "Recorded here so the table is not mistaken for the whole picture:\n",
        "* **Research.** The majority of the project's research was done by **Mahesh "
        "Rajendra** — domain reading on loan origination and underwriting, lending-policy "
        "structure, agentic architecture and evaluation approaches. None of it produces "
        "commits.\n"
        "* **Collaborative development.** The two members worked together rather than "
        "splitting the project into independently committed halves, so commit counts do "
        "not reflect the split of effort.\n"
        "* **Internal review.** Peer review of code and project work — see "
        "[PEER_REVIEW.md](PEER_REVIEW.md).\n"
        "* **Pre-commit work.** Design, debugging and iteration before each commit.\n",
        "\n## Engagement attestation (REQ-008)\n",
        f"> {ATTESTATION}\n",
        "Each member confirms for themselves only. Nobody signs on anyone else's behalf.\n",
        "### Krishna Annavaram\n",
        "```\nConfirmation: __________________________\nDate:         __________________________\n```\n",
        "### Mahesh Rajendra\n",
        "```\nConfirmation: __________________________\nDate:         __________________________\n```\n",
        "**Status: PENDING — unsigned.**\n",
        "\n### After both blocks are signed\n",
        "Replace the `REQ-008-T01` entry in "
        "`requirements_validation_tests/manual_evidence/manual_attestations.json` with:\n",
        "```json",
        '"REQ-008-T01": {',
        '  "requirement_id": "REQ-008",',
        '  "what_must_be_attested": "The engagement ran for the stated duration.",',
        '  "source_location": "Section 2. Engagement Overview - table row 1",',
        '  "source_text": "Duration | 20 hours",',
        '  "what_would_count_as_evidence": "The actual elapsed engagement time against the stated 20 hours.",',
        '  "where_to_look": "docs/team/WORKLOG.md and reports/team_worklog.xlsx.",',
        '  "evidence": "docs/team/WORKLOG.md and reports/team_worklog.xlsx. Combined development, testing, evaluation, review, research and documentation activity met or exceeded 20 hours. '
        f'Git shows a {window_h}-hour wall-clock window over {len(commits)} commits; exact task-level hours are team-attested rather than mechanically derivable, and the research effort is not represented in commits at all.",',
        '  "attested_by": "Krishna Annavaram; Mahesh Rajendra",',
        '  "attested_on": "YYYY-MM-DD"',
        "}",
        "```\n",
        "Set `attested_on` to the date the **second** signature was added, then re-run "
        "the validator and `build_gap_analysis.py`. `REQ-008` moves to PASS at that "
        "point and not before.\n",
        "\n## Evidence precedence\n",
        "In descending order of authority, as used above:\n",
        "1. Git timestamps\n2. Test and evaluation timestamps\n3. Trace and report "
        "timestamps\n4. This worklog\n5. The signed team attestation\n",
        "Conversations with any AI assistant are **not** evidence of elapsed hours and "
        "are not cited here. They can help reconstruct what was worked on; they prove "
        "nothing about duration.\n",
    ]
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(out) + "\n", encoding="utf-8")


def main() -> int:
    from src.console import use_utf8_stdio

    use_utf8_stdio()

    commits = load_commits()
    if not commits:
        print("No commits found.")
        return 2
    sessions = build_sessions(commits)
    artifacts = evidence_artifacts()

    write_xlsx(commits, sessions, artifacts, XLSX_OUT)
    write_markdown(commits, sessions, artifacts, MD_OUT)

    derivable = [s for s in sessions if s["span_hours"] is not None]
    window_h = round((commits[-1]["when"] - commits[0]["when"]).total_seconds() / 3600.0, 1)
    print(f"commits            : {len(commits)}")
    print(f"sessions           : {len(sessions)} (gap {SESSION_GAP_HOURS}h)")
    print(f"elapsed wall-clock : {window_h} h  (window, NOT effort)")
    print(f"observed span      : {round(sum(s['span_hours'] for s in derivable), 2)} h "
          f"across {len(derivable)} sessions (lower bound)")
    print(f"evidence artifacts : {len(artifacts)}")
    print(f"\nWritten:\n  {XLSX_OUT}\n  {MD_OUT}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
